"""DOCX text extraction."""
import logging
from typing import Optional

logger = logging.getLogger(__name__)


def extract_text_from_docx(docx_bytes: bytes) -> Optional[str]:
    """Extract text from DOCX file.
    
    Args:
        docx_bytes: DOCX file content as bytes
        
    Returns:
        Extracted text or None if extraction failed
    """
    try:
        import io
        from docx import Document
        
        doc = Document(io.BytesIO(docx_bytes))
        text_parts = []
        
        # Extract text from paragraphs
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text)
        
        # Extract text from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text_parts.append(cell.text)
        
        text = "\n".join(text_parts)
        if text.strip():
            logger.info(f"Extracted {len(text)} characters from DOCX")
            return text
        
        return None
    except Exception as e:
        logger.error(f"DOCX extraction failed: {e}")
        return None
